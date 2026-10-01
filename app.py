"""
Automation Assistant - Streamlit app for the hackathon.

Two tools in one. Employees choose what they need:
  1. Create an Excel macro (VBA agent)
  2. Write a process document (process document agent, no macro needed)

Purpose of option 1: help employees with little time or technical knowledge create a working Excel
automation macro. They upload a file, describe the task in plain English, and receive:
  - a ready-to-import macro (.bas) and a data workbook to import it into,
  - a plain-English explanation of what the macro will do and change,
  - a way to ask for changes or paste an error message to get a fixed version,
  - a process document (Word) for the team: manual steps, what the macro automates, items to confirm.

Run:
    pip install streamlit pandas openpyxl anthropic openai python-docx
    streamlit run app.py

Add your key through the sidebar, Streamlit secrets, or an environment variable.
Keys starting with "sk-ant-" use Claude (Anthropic); other keys use OpenAI.
"""

import inspect
import io
import json
import os
import re
import zipfile
from datetime import date

import pandas as pd
import streamlit as st

# --------------------------------------------------------------------------- #
# Page setup
# --------------------------------------------------------------------------- #
st.set_page_config(page_title="Workflow Studio", page_icon="📊", layout="wide")

st.markdown(
    """
    <style>
        .block-container { padding-top: 2rem; max-width: 1200px; }
        h1 { letter-spacing: -0.02em; }
        div.stButton > button[kind="primary"] { width: 100%; font-weight: 600; padding: 0.6rem 0; }
    </style>
    """,
    unsafe_allow_html=True,
)

# Keep this empty in source control. Use the sidebar, Streamlit secrets, or an environment variable.
HARDCODED_API_KEY = ""

# Default model per provider. Change it in the sidebar if your account uses a different one.
DEFAULT_MODELS = {"anthropic": "claude-sonnet-5-5", "openai": "gpt-5.5"}
LANGUAGES = ["English", "Japanese", "Spanish", "French", "German", "Portuguese", "Chinese"]
MODULE_NAME = "GeneratedMacro"
MODE_MACRO = "Create an Excel macro"
MODE_DOC = "Write a process document"
MODE_HELP = {
    MODE_MACRO: "Turn a task into a ready-to-run Excel macro, with a plain-English explanation.",
    MODE_DOC: "Turn your notes into a clear, step-by-step process document your team can follow.",
}
MAX_TEXT_CHARS = 20_000
CODE_BOX_HEIGHT = 450  # fixed height (pixels) of the code box; longer code scrolls inside it

# Things worth a plain-language heads-up before the macro is run.
RISKY_PATTERNS = {
    r"\bKill\b": "deletes files from your computer",
    r"\bRmDir\b": "removes folders from your computer",
    r"\bShell\b|WScript\.Shell": "starts other programs on your computer",
    r"\bFileSystemObject\b": "reads or writes files outside Excel",
    r"\bURLDownloadToFile\b": "downloads files from the internet",
    r"\bXMLHTTP\b|\bWinHttp\b": "connects to the internet",
    r"\bSendKeys\b": "simulates keyboard typing, which can be unreliable",
    r"\.Delete\b": "deletes rows, columns or sheets in your workbook",
    r"\.ClearContents\b|\.Clear\b": "erases cell contents",
    r"\.RemoveDuplicates\b": "removes duplicate rows",
}


# --------------------------------------------------------------------------- #
# File helpers
# --------------------------------------------------------------------------- #
def excel_col_letter(n):
    letters = ""
    while n:
        n, r = divmod(n - 1, 26)
        letters = chr(65 + r) + letters
    return letters


def read_upload(uploaded):
    """Return (DataFrame | None, raw_text | None, sheet_name)."""
    name = uploaded.name.lower()
    data = uploaded.getvalue()

    if name.endswith(".csv"):
        return pd.read_csv(io.BytesIO(data)), None, "Data"
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        xls = pd.ExcelFile(io.BytesIO(data))
        first = xls.sheet_names[0]
        return xls.parse(first), None, first
    if name.endswith(".json"):
        try:
            return pd.json_normalize(json.loads(data)), None, "Data"
        except Exception:
            return None, data.decode("utf-8", errors="replace"), "Data"
    return None, data.decode("utf-8", errors="replace"), "Data"


def describe_file(name, df, text, sheet_name, include_samples=True):
    """Compact description of the file for the model. Sample rows are optional (privacy mode)."""
    if df is not None:
        cols = [
            f"- Column {excel_col_letter(i + 1)}: {col} ({dtype})"
            for i, (col, dtype) in enumerate(df.dtypes.astype(str).items())
        ]
        description = (
            f"Source file: {name}\n"
            f"The data will be on a worksheet named \"{sheet_name}\", headers in row 1, "
            f"data from row 2 down to row {len(df) + 1}.\n\n"
            "Columns:\n" + "\n".join(cols)
        )
        if include_samples:
            description += f"\n\nFirst 5 rows (CSV):\n{df.head(5).to_csv(index=False)}"
        else:
            description += "\n\n(Sample rows were withheld for privacy.)"
        return description
    truncated = text[:MAX_TEXT_CHARS]
    note = "\n[truncated]" if len(text) > MAX_TEXT_CHARS else ""
    return f"Source file: {name}\n\nContent:\n{truncated}{note}"


def build_workbook(df, sheet_name):
    """Data workbook (.xlsx) so the user can import the .bas into it and save as .xlsm."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    return buf.getvalue()


def suggest_examples(df):
    """Clickable example requests built from the user's own column names."""
    cat = next((c for c in df.columns if pd.api.types.is_string_dtype(df[c]) and 2 <= df[c].nunique() <= 20), None)
    num = next((c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])), None)

    examples = [("Format the header row",
                 "Make the header row bold with a light blue background, freeze it, and fit all column widths.")]
    if cat and num:
        examples += [
            (f"Summary by {cat}", f"Create a new sheet that shows the total and average {num} for each {cat}."),
            (f"One sheet per {cat}", f"Split the data into a separate sheet for each {cat}, keeping the header row."),
            (f"Highlight low {num}", f"Highlight in red every row where {num} is below the average {num}."),
        ]
    examples.append(("Clean the data", "Remove duplicate rows and trim extra spaces in all text cells."))
    return examples


# --------------------------------------------------------------------------- #
# AI calls
# --------------------------------------------------------------------------- #
SYSTEM_PROMPT = """You are an expert Excel VBA developer. Write a complete, working VBA module \
based on the user's data description and instructions.

Rules:
- Return ONLY VBA code, no explanations and no markdown fences.
- Start with Option Explicit and declare every variable with Dim.
- Put the main entry point in a Public Sub with a clear name and no required arguments, \
so it can be run from Alt+F8.
- Work on ThisWorkbook and refer to worksheets by the exact name given. Never rely on ActiveSheet.
- Find the last row/column dynamically (for example with .Cells(.Rows.Count, 1).End(xlUp).Row). \
Do not hard-code row counts.
- Add error handling (On Error GoTo ...) that restores Application.ScreenUpdating and \
Application.Calculation and shows a MsgBox with the error description.
- Turn off ScreenUpdating while the macro runs and turn it back on at the end.
- Use only built-in Excel/VBA objects. No external references, no late-bound COM objects, \
no file deletion, no Shell calls, no network access.
- Add short comments explaining each step.
- End with a MsgBox that tells the user the macro has finished."""

EXPLAIN_PROMPT = """You explain Excel macros to office workers with no programming knowledge. \
Write in {language}. Use short sentences and everyday words. Never show code, and avoid technical \
terms (if you must use one, explain it in a few words).

You will receive a VBA macro. Use exactly these four Markdown sections:

### What this macro does
Two or three sentences.

### What happens when you run it
A numbered list with one plain step per line, in the order the macro does them.

### What will change in your workbook
A bullet list of the sheets, columns, colors or cells that are added, changed or erased. \
Say clearly if anything is deleted or overwritten, and say so if nothing is.

### Before you run it
Two or three short bullets of practical advice, such as working on a copy and which sheet name \
the macro expects."""


PROCESS_DOC_PROMPT = """You are a technical writer who documents business processes in plain, friendly \
language for HR, Compensation, and data analytics teams.

You will receive the user's description of what they wanted automated, details about the workbook \
(sheet and column names), the name of the VBA macro that automates it, and the macro's code. \
Your document will sit alongside the macro so that anyone on the team can understand what it does \
and how to do the work by hand if needed.

YOUR JOB

1. Read the whole input and understand what the process does, who does it, what goes in, and what \
comes out.
2. Rewrite it so it is clear, short, and easy to follow for someone doing it for the first time.
3. Treat the macro code as the source of truth for what the macro does. Keep the order and logic of \
the macro's steps in your manual steps. Do not remove, merge, or reorder steps in a way that changes \
how the process works.
4. If the macro does something the user's description did not mention, or the description asks for \
something the macro does not do, say so in section 6.
5. If something is missing or unclear, you may fill the gap based on common sense or best practice, \
but mark every such addition with: (Assumed – please confirm)
6. Never add new logic silently.

STYLE

- Write the whole document in {language}. Keep the marker text (Assumed – please confirm) exactly as \
written, in English, even when you write in another language, so it can be found easily.
- Plain, short sentences. One action per step.
- Write to the reader directly ("Open the file", not "The analyst opens the file").
- Say which role does a step only when it changes hands.
- No jargon unless the input uses it; if it does, explain it the first time it appears.
- Avoid filler, repetition, and overly formal phrasing.
- Never show VBA code in the document.

OUTPUT FORMAT

Produce the document in Markdown, with no code fences around it. Use exactly these six sections, \
in this order, each as a level-2 heading with its number (translate the heading text into the \
document language). Use numbered lists for steps; use a table only in section 5.

## 1. What this process does
2-3 sentences: what it is for and what the end result is.

## 2. Before you start
Bullet list of what you need: files, sheets, tools, access, and any conditions that must be true \
(for example, "The workbook has a sheet named Employees").

## 3. Steps (manual)
Numbered list. Each step is one clear action. Add the input or output for a step only when it is \
not obvious. Mark any inferred step with (Assumed – please confirm).

## 4. Result
1-2 sentences describing what the finished output looks like and how to check it worked.

## 5. What the macro automates
A two-column Markdown table with the header "Manual step(s)" | "What the macro does". Map each \
manual step (by number) to the macro's behavior. Below the table, note anything the macro does NOT \
do that still has to be done by hand.

## 6. Notes and assumptions
One list that collects every (Assumed – please confirm) item, plus any difference you found between \
the user's description and what the macro actually does, so the process owner can review everything \
in one place. If there are none, write "None".

FINAL CHECK

- Sections appear exactly as listed above.
- The macro's step order and logic are preserved.
- Every inferred addition is marked (Assumed – please confirm) and repeated in section 6.
- A new team member could follow section 3 without asking questions."""


def standalone_doc_prompt(language, with_macro, has_code):
    """Prompt for the stand-alone process document writer. Sections adapt to whether a macro exists."""
    intro = (
        "You are a technical writer who documents business processes in plain, friendly language for HR, "
        "Compensation, and data analytics teams.\n\n"
        "The user will give you a description of a manual process (notes, a rough draft, or an existing document)."
    )
    if with_macro:
        intro += (" A VBA macro automates this process, and your document will sit alongside the macro so that "
                  "anyone on the team can understand what it does and how to do the work by hand if needed.")
    else:
        intro += " Your document will help anyone on the team understand the process and do the work by hand."

    job = [
        "Read the whole input and understand what the process does, who does it, what goes in, and what comes out.",
        "Rewrite it so it is clear, short, and easy to follow for someone doing it for the first time.",
        "Keep the original order and logic of the steps. Do not remove, merge, or reorder steps in a way that "
        "changes how the process works.",
        "If something is missing or unclear, you may fill the gap based on common sense or best practice, "
        "but mark every such addition with: (Assumed – please confirm)",
        "Never add new logic silently.",
    ]
    if with_macro and has_code:
        job.append("Use the macro code to describe what the macro does in the macro section. If the macro does "
                   "something the description did not mention, or the description asks for something the macro "
                   "does not do, say so in the notes section.")
    elif with_macro:
        job.append("You only know the macro's name, not its code. In the macro section, describe what the macro "
                   "does only where the process description says so, and mark any guess with "
                   "(Assumed – please confirm).")

    style = [
        f"Write the whole document in {language}. Keep the marker text (Assumed – please confirm) exactly as "
        "written, in English, even when you write in another language, so it can be found easily.",
        "Plain, short sentences. One action per step.",
        'Write to the reader directly ("Open the file", not "The analyst opens the file").',
        "Say which role does a step only when it changes hands.",
        "No jargon unless the input uses it; if it does, explain it the first time it appears.",
        "Avoid filler, repetition, and overly formal phrasing.",
    ]
    if with_macro:
        style.append("Never show VBA code in the document.")

    sections = [
        ("What this process does", "2-3 sentences: what it is for and what the end result is."),
        ("Before you start", "Bullet list of what you need: files, sheets, tools, access, and any conditions that "
                             'must be true (for example, "The workbook has a sheet named Employees").'),
        ("Steps (manual)", "Numbered list. Each step is one clear action. Add the input or output for a step only "
                           "when it is not obvious. Mark any inferred step with (Assumed – please confirm)."),
        ("Result", "1-2 sentences describing what the finished output looks like and how to check it worked."),
    ]
    if with_macro:
        sections.append(("What the macro automates",
                         'A two-column Markdown table with the header "Manual step(s)" | "What the macro does". '
                         "Map each manual step (by number) to the macro's behavior. Below the table, note anything "
                         "the macro does NOT do that still has to be done by hand."))
    notes = ("One list that collects every (Assumed – please confirm) item so the process owner can review them "
             "in one place.")
    if with_macro and has_code:
        notes += " Also list any difference you found between the description and what the macro actually does."
    notes += ' If there are none, write "None".'
    sections.append(("Notes and assumptions", notes))

    table_rule = "use a table only in section 5" if with_macro else "do not use tables"
    body = "\n\n".join(f"## {i}. {title}\n{text}" for i, (title, text) in enumerate(sections, 1))
    final_check = [
        "Sections appear exactly as listed above.",
        "Original step order and logic are preserved.",
        f"Every inferred addition is marked (Assumed – please confirm) and repeated in section {len(sections)}.",
        "A new team member could follow the steps without asking questions.",
    ]
    return (
        f"{intro}\n\nYOUR JOB\n\n" + "\n".join(f"{i}. {t}" for i, t in enumerate(job, 1))
        + "\n\nSTYLE\n\n" + "\n".join(f"- {t}" for t in style)
        + f"\n\nOUTPUT FORMAT\n\nProduce the document in Markdown, with no code fences around it. Use exactly these "
          f"{len(sections)} sections, in this order, each as a level-2 heading with its number (translate the "
          f"heading text into the document language). Use numbered lists for steps; {table_rule}.\n\n{body}"
        + "\n\nFINAL CHECK\n\n" + "\n".join(f"- {t}" for t in final_check)
    )


def detect_provider(api_key):
    """Anthropic keys start with 'sk-ant-'; anything else is treated as an OpenAI key."""
    return "anthropic" if api_key.strip().startswith("sk-ant-") else "openai"


def call_llm(system, user, api_key, model, provider, max_tokens=8000):
    if provider == "anthropic":
        import anthropic

        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model=model,
            max_tokens=max_tokens,
            system=system,
            messages=[{"role": "user", "content": user}],
        )
        return "".join(b.text for b in response.content if b.type == "text")

    from openai import OpenAI

    client = OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model=model,
        # Reasoning models count their thinking tokens in this limit, so keep it generous.
        max_completion_tokens=max_tokens,
        messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
    )
    return response.choices[0].message.content or ""


def generate_vba(file_context, instructions, api_key, model, provider):
    user = f"{file_context}\n\n---\nInstructions:\n{instructions}"
    return call_llm(SYSTEM_PROMPT, user, api_key, model, provider)


def revise_vba(file_context, current_code, change_request, api_key, model, provider):
    user = (
        f"{file_context}\n\n---\nHere is the current macro:\n{current_code}\n\n---\n"
        f"The user wants this change, or ran into this problem:\n{change_request}\n\n"
        "Return the complete updated macro."
    )
    return call_llm(SYSTEM_PROMPT, user, api_key, model, provider)


def explain_vba(code, language, api_key, model, provider):
    system = EXPLAIN_PROMPT.replace("{language}", language)
    return call_llm(system, code, api_key, model, provider, max_tokens=4000).strip()


def generate_process_doc(requests, process_name, macro_name, sheet_name, columns, code,
                         language, api_key, model, provider):
    """Write the team process document. Only column names and the macro code are sent, never data rows."""
    described = []
    for i, request in enumerate(requests):
        label = "Original request" if i == 0 else f"Later change or problem reported (#{i})"
        described.append(f"{label}: {request}")
    column_list = ", ".join(f"{excel_col_letter(i + 1)}: {c}" for i, c in enumerate(columns)) or "Not available"

    user = (
        "Process description:\n" + "\n".join(described) + "\n\n"
        f"Process name (optional): {process_name or 'Not provided - choose a short name from the description'}\n"
        f"Macro name: {macro_name}\n"
        f"Worksheet: \"{sheet_name}\" (headers in row 1, data from row 2)\n"
        f"Columns: {column_list}\n\n"
        f"Macro code:\n{code}"
    )
    system = PROCESS_DOC_PROMPT.replace("{language}", language)
    text = call_llm(system, user, api_key, model, provider).strip()
    # Remove code fences if the model wrapped the document in them.
    fence = re.match(r"^```(?:markdown|md)?\s*\n(.*?)\n```\s*$", text, re.DOTALL | re.IGNORECASE)
    return fence.group(1).strip() if fence else text


def generate_standalone_doc(description, process_name, with_macro, macro_name, macro_code, language,
                            api_key, model, provider, previous_doc=None, change_request=None):
    """Write (or revise) a process document from the user's own description. Used by the Process document writer."""
    truncated = description[:MAX_TEXT_CHARS]
    if len(description) > MAX_TEXT_CHARS:
        truncated += "\n[truncated]"
    parts = [
        f"Process description:\n{truncated}",
        f"Process name (optional): {process_name or 'Not provided - choose a short name from the description'}",
    ]
    if with_macro:
        parts.append(f"Macro name: {macro_name or 'Not provided'}")
        if macro_code.strip():
            parts.append(f"Macro code:\n{macro_code.strip()[:MAX_TEXT_CHARS]}")
    if previous_doc:
        parts.append(f"Here is the current version of the document:\n{previous_doc}\n\n"
                     f"The user wants this change:\n{change_request}\n\n"
                     "Return the complete updated document in the same format.")
    system = standalone_doc_prompt(language, with_macro, bool(macro_code.strip()))
    text = call_llm(system, "\n\n".join(parts), api_key, model, provider).strip()
    fence = re.match(r"^```(?:markdown|md)?\s*\n(.*?)\n```\s*$", text, re.DOTALL | re.IGNORECASE)
    return fence.group(1).strip() if fence else text


def read_notes_file(uploaded_file):
    """Text from an uploaded notes file (.docx, .txt or .md). Reading .docx needs: pip install python-docx"""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if not name.endswith(".docx"):
        return data.decode("utf-8", errors="replace")

    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    lines = []
    for child in document.element.body.iterchildren():  # keeps paragraphs and tables in document order
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text
            if text.strip():
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def list_models(api_key, provider):
    """Return the model IDs this API key can use."""
    if provider == "anthropic":
        import anthropic

        return sorted(m.id for m in anthropic.Anthropic(api_key=api_key).models.list(limit=100))

    from openai import OpenAI

    ids = sorted(m.id for m in OpenAI(api_key=api_key).models.list())
    return [i for i in ids if i.startswith(("gpt-", "chat-latest", "o"))]


def friendly_error(exc):
    """Return (plain-language message, technical detail)."""
    raw = str(exc)
    low = raw.lower()
    if any(k in low for k in ("401", "authentication", "invalid x-api-key", "incorrect api key", "invalid api key")):
        return "The API key was rejected. Check that it is correct and still active.", raw
    if any(k in low for k in ("404", "not_found", "model_not_found", "does not exist")):
        return ("The model name was not found for this key. Open Connection settings in the sidebar, "
                "click Check available models, and use a name from the list."), raw
    if any(k in low for k in ("429", "rate limit", "rate_limit", "quota", "insufficient", "credit balance", "billing")):
        return "The API account has reached a usage limit or has no credits left. Check billing with your provider.", raw
    return "Something went wrong while contacting the AI service. Please try again.", raw


# --------------------------------------------------------------------------- #
# VBA post-processing
# --------------------------------------------------------------------------- #
def clean_vba(raw):
    """Strip markdown fences, normalize to Option Explicit, drop any module attribute line."""
    text = raw.strip()
    fence = re.search(r"```(?:vba|vb|basic)?\s*\n(.*?)```", text, re.DOTALL | re.IGNORECASE)
    if fence:
        text = fence.group(1).strip()
    text = re.sub(r"^\s*Attribute VB_Name\s*=.*\n?", "", text, flags=re.IGNORECASE | re.MULTILINE)
    if not re.search(r"^\s*Option Explicit", text, re.IGNORECASE | re.MULTILINE):
        text = "Option Explicit\n\n" + text
    return text.strip()


def to_bas_bytes(code):
    """Importable .bas file: module header, CRLF line endings, ANSI encoding (what the VBE expects)."""
    body = f'Attribute VB_Name = "{MODULE_NAME}"\n{code}\n'
    body = body.replace("\r\n", "\n").replace("\n", "\r\n")
    return body.encode("cp1252", errors="replace")


def show_code_box(code):
    """Show the VBA in a box with a fixed height; long code scrolls inside the box."""
    if "height" in inspect.signature(st.code).parameters:  # newer Streamlit
        st.code(code, language="vbnet", line_numbers=True, height=CODE_BOX_HEIGHT)
    else:  # older Streamlit: use a fixed-height container instead
        with st.container(height=CODE_BOX_HEIGHT):
            st.code(code, language="vbnet", line_numbers=True)


def main_macro_name(code):
    match = re.search(r"^\s*(?:Public\s+)?Sub\s+(\w+)\s*\(\s*\)", code, re.IGNORECASE | re.MULTILINE)
    return match.group(1) if match else None


def review_vba(code):
    """Return plain-language warnings about the generated code."""
    warnings = []
    if main_macro_name(code) is None:
        warnings.append("This macro may not appear in the Alt+F8 list. Ask for a change: "
                        "\"Make the macro runnable from the macro list.\"")
    for pattern, message in RISKY_PATTERNS.items():
        if re.search(pattern, code, re.IGNORECASE):
            warnings.append(f"Heads-up: this macro {message}. Make sure that is what you want, "
                            "and run it on a copy of your file first.")
    return warnings


# --------------------------------------------------------------------------- #
# Process document helpers
# --------------------------------------------------------------------------- #
def count_assumptions(markdown_text):
    """Count '(Assumed - please confirm)' marks in the steps (before the final notes section repeats them)."""
    headings = list(re.finditer(r"^##\s+", markdown_text, flags=re.MULTILINE))
    body = markdown_text[: headings[-1].start()] if len(headings) > 1 else markdown_text
    return body.count("Assumed")


def _add_inline(paragraph, text):
    """Add text to a paragraph, turning **bold** and `code` into formatted runs."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)


def markdown_to_docx(markdown_text, title):
    """Convert the process document (simple Markdown) to a Word file. Needs: pip install python-docx"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            style.element.rPr.rFonts.attrib.pop(qn(attr), None)

    doc.add_heading(title or "Process document", 0)

    lines = markdown_text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # Table: consecutive lines that start with "|"
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip().replace("<br>", " ") for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):  # skip the |---|---| row
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                table.autofit = False
                # Two-column "step | what the macro does" tables: narrow first column.
                widths = [1.3, 5.2] if ncols == 2 else [6.5 / ncols] * ncols
                for col_idx, column in enumerate(table.columns):
                    column.width = Inches(widths[col_idx])
                for r_idx, row in enumerate(rows):
                    for c_idx in range(ncols):
                        cell = table.cell(r_idx, c_idx)
                        cell.width = Inches(widths[c_idx])
                        paragraph = cell.paragraphs[0]
                        _add_inline(paragraph, row[c_idx] if c_idx < len(row) else "")
                        if r_idx == 0:
                            for run in paragraph.runs:
                                run.bold = True
                            shade = OxmlElement("w:shd")
                            shade.set(qn("w:val"), "clear")
                            shade.set(qn("w:color"), "auto")
                            shade.set(qn("w:fill"), "DDEBF7")
                            cell._tc.get_or_add_tcPr().append(shade)
                doc.add_paragraph()
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", stripped)
        numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", line)
        bullet = re.match(r"^(\s*)[-*•]\s+(.*)", line)

        if heading:
            doc.add_heading(heading.group(2).replace("**", ""), min(len(heading.group(1)), 3))
        elif numbered:
            paragraph = doc.add_paragraph()
            fmt = paragraph.paragraph_format
            fmt.left_indent = Inches(0.35 + 0.3 * (len(numbered.group(1)) // 3))
            fmt.first_line_indent = Inches(-0.3)
            fmt.space_after = Pt(3)
            fmt.tab_stops.add_tab_stop(fmt.left_indent)
            _add_inline(paragraph, f"{numbered.group(2)}.\t{numbered.group(3)}")
        elif bullet:
            nested = len(bullet.group(1)) >= 2
            paragraph = doc.add_paragraph(style="List Bullet 2" if nested else "List Bullet")
            _add_inline(paragraph, bullet.group(2))
        elif stripped and not re.fullmatch(r"-{3,}", stripped):
            _add_inline(doc.add_paragraph(), stripped)
        i += 1

    note = doc.add_paragraph()
    run = note.add_run(
        f"Created by Excel Automation Assistant on {date.today():%Y-%m-%d}. "
        "Please review the notes and assumptions before sharing this document."
    )
    run.italic = True
    run.font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_package(code, doc_bytes, doc_markdown, workbook_bytes):
    """Zip with the macro, the process document and (if available) the data workbook."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{MODULE_NAME}.bas", to_bas_bytes(code))
        if doc_bytes:
            z.writestr("ProcessDocument.docx", doc_bytes)
        else:
            z.writestr("ProcessDocument.md", doc_markdown)
        if workbook_bytes:
            z.writestr("data_workbook.xlsx", workbook_bytes)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Sidebar
# --------------------------------------------------------------------------- #
with st.sidebar:
    st.header("Settings")
    language = st.selectbox("Language (explanations and documents)", LANGUAGES)
    privacy_mode = st.checkbox(
        "Privacy mode",
        help="Send only column names to the AI service. Your data rows are not shared.",
    )

    def read_secret(name):
        """Read a key from .streamlit/secrets.toml (or the cloud Secrets box). Empty if there is none."""
        try:
            return str(st.secrets.get(name, "")).strip()
        except Exception:
            return ""

    default_key = (
        HARDCODED_API_KEY.strip()
        or read_secret("ANTHROPIC_API_KEY")
        or read_secret("OPENAI_API_KEY")
        or os.environ.get("ANTHROPIC_API_KEY", "")
        or os.environ.get("OPENAI_API_KEY", "")
    )
    with st.expander("Connection settings", expanded=not default_key):
        api_key = st.text_input(
            "API key (Anthropic or OpenAI)",
            value=default_key,
            type="password",
            help="Paste your key here, or use Streamlit secrets or an environment variable.",
        )
        provider = detect_provider(api_key)
        model = st.text_input("Model", value=DEFAULT_MODELS[provider])

        if st.button("Check available models"):
            if not api_key.strip():
                st.warning("Add your API key first.")
            else:
                try:
                    st.code("\n".join(list_models(api_key.strip(), provider)), language="text")
                except Exception as exc:
                    st.error(friendly_error(exc)[0])

    if api_key.strip():
        provider_name = "Anthropic (Claude)" if provider == "anthropic" else "OpenAI"
        st.success(f"{provider_name} key loaded (ends with ...{api_key.strip()[-4:]}).")
    else:
        st.error("No API key found. Paste it in Connection settings, or use Streamlit secrets or an environment variable.")

# --------------------------------------------------------------------------- #
# Main page
# --------------------------------------------------------------------------- #
st.title("Automation Assistant")
st.caption("Two tools in one. Choose what you need today.")

mode = st.radio("What do you need?", [MODE_MACRO, MODE_DOC], horizontal=True, key="mode",
                label_visibility="collapsed")
st.caption(MODE_HELP[mode])

# Streamlit forgets the text of widgets that are hidden. Re-saving it keeps what the user typed
# when they switch between the two options.
for _key in ("instructions", "process_name", "pd_description", "pd_name", "pd_has_macro",
             "pd_macro_name", "pd_macro_code"):
    if _key in st.session_state:
        st.session_state[_key] = st.session_state[_key]


# --------------------------------------------------------------------------- #
# Option 1: Excel macro assistant
# --------------------------------------------------------------------------- #
def render_macro_assistant():
    left, right = st.columns(2, gap="large")


    def set_instructions(prompt):
        st.session_state["instructions"] = prompt


    def produce(mode, change_request=""):
        """Generate (mode='new') or update (mode='update') the macro, then explain it. Returns True on success."""
        context = describe_file(uploaded.name, df, text, sheet_name, include_samples=not privacy_mode)
        key, mdl = api_key.strip(), model.strip()

        with st.status("Working on your macro...", expanded=True) as status:
            try:
                if mode == "new":
                    st.write("Reading your file and instructions...")
                    st.write("Writing the macro...")
                    raw = generate_vba(context, instructions, key, mdl, provider)
                else:
                    st.write("Updating the macro with your changes...")
                    raw = revise_vba(context, st.session_state["code"], change_request, key, mdl, provider)

                if not raw.strip():
                    raise ValueError("The AI returned an empty answer. Please try again.")
                code = clean_vba(raw)

                st.write("Explaining it in plain English...")
                try:
                    explanation = explain_vba(code, language, key, mdl, provider)
                except Exception:
                    explanation = None  # the macro is still usable without the explanation

                st.session_state["code"] = code
                if mode == "new":
                    st.session_state["requests"] = [instructions]
                    st.session_state.pop("process_doc", None)  # a brand-new macro needs a new document
                else:
                    st.session_state.setdefault("requests", []).append(change_request)
                st.session_state["columns"] = list(df.columns) if df is not None else []
                st.session_state["explanation"] = explanation
                st.session_state["workbook"] = build_workbook(df, sheet_name) if df is not None else None
                st.session_state["sheet_name"] = sheet_name
                st.session_state["version"] = st.session_state.get("version", 0) + 1 if mode == "update" else 1
                status.update(label="Your macro is ready.", state="complete", expanded=False)
                return True
            except Exception as exc:
                status.update(label="Something went wrong.", state="error", expanded=True)
                friendly, detail = friendly_error(exc)
                st.error(friendly)
                with st.expander("Technical details"):
                    st.code(detail, language="text")
                return False


    with left:
        st.subheader("1. Upload your file")
        uploaded = st.file_uploader(
            "Drag a file here or browse",
            type=["csv", "xlsx", "xlsm", "xls", "json", "txt"],
            label_visibility="collapsed",
        )
        st.caption(
            "Only your column names and the first 5 rows are sent to the AI service. "
            "Turn on Privacy mode in the sidebar to send column names only."
        )

        df, text, sheet_name = None, None, "Data"
        if uploaded is not None:
            try:
                df, text, sheet_name = read_upload(uploaded)
            except Exception as exc:
                st.error(f"Could not read {uploaded.name}: {exc}")
                uploaded = None

        if uploaded is not None:
            # A new file means the previous macro no longer applies.
            if st.session_state.get("source_name") != uploaded.name:
                for k in ("code", "explanation", "workbook", "version", "sheet_name", "requests", "columns",
                          "process_doc", "process_doc_version"):
                    st.session_state.pop(k, None)
                st.session_state["source_name"] = uploaded.name

            if df is not None:
                st.success(f"{uploaded.name}: {len(df):,} rows, {df.shape[1]} columns (sheet \"{sheet_name}\")")
                st.dataframe(df.head(10), use_container_width=True)
                with st.expander("Column names you can mention in your instructions"):
                    st.markdown(", ".join(f"`{c}`" for c in df.columns))
            else:
                st.success(f"{uploaded.name} loaded")
                st.code(text[:1500] + ("\n..." if len(text) > 1500 else ""), language="text")

        st.subheader("2. Say what you want")
        hint = st.session_state.pop("macro_hint", None)
        if hint:
            st.info(hint)
        instructions = st.text_area(
            "Instructions",
            key="instructions",
            height=150,
            placeholder="Example: Highlight every row where Sales is below 1,000 in red, "
            "and add a summary sheet with total sales by region.",
            label_visibility="collapsed",
        )

        if df is not None:
            st.caption("Need ideas? Click one to fill the box:")
            examples = suggest_examples(df)
            cols = st.columns(2)
            for i, (label, prompt) in enumerate(examples):
                cols[i % 2].button(label, key=f"example_{i}", on_click=set_instructions, args=(prompt,),
                                   use_container_width=True)

        generate = st.button("Create my macro", type="primary")

    with right:
        st.subheader("3. Your macro")

        if generate:
            if uploaded is None:
                st.warning("Upload a file first.")
            elif not instructions.strip():
                st.warning("Tell me what you want the macro to do first.")
            elif not api_key.strip():
                st.error("No API key found, so nothing was created. "
                         "Paste it in Connection settings, or use Streamlit secrets or an environment variable.")
            else:
                produce("new")

        if "code" in st.session_state:
            code = st.session_state["code"]
            explanation = st.session_state.get("explanation")
            version = st.session_state.get("version", 1)

            flash = st.session_state.pop("flash", None)
            if flash:
                st.success(flash)
            st.caption(f"Version {version}")

            for warning in review_vba(code):
                st.warning(warning)

            d1, d2 = st.columns(2)
            d1.download_button(
                "Download macro (.bas)",
                data=to_bas_bytes(code),
                file_name=f"{MODULE_NAME}.bas",
                mime="text/plain",
                use_container_width=True,
                type="primary",
            )
            if st.session_state.get("workbook"):
                d2.download_button(
                    "Download data workbook (.xlsx)",
                    data=st.session_state["workbook"],
                    file_name="data_workbook.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

            tab_explain, tab_howto, tab_doc, tab_code = st.tabs(
                ["What it does", "How to use it", "Process document", "The code (advanced)"]
            )

            with tab_explain:
                if explanation:
                    st.markdown(explanation)
                else:
                    st.info("The explanation could not be created this time. The macro itself is ready; "
                            "click Update the macro below to try again.")

            with tab_howto:
                macro_name = main_macro_name(code) or "your macro"
                expected_sheet = st.session_state.get("sheet_name", "Data")
                st.markdown(
                    f"""
    **Before you start:** make a copy of your workbook. Some macro changes cannot be undone.

    1. Open your Excel workbook (or the data workbook you downloaded above).
    2. Press **Alt + F11** to open the macro editor.
    3. Choose **File > Import File...** and select `{MODULE_NAME}.bas`.
    4. Close the editor window.
    5. Press **Alt + F8**, select **{macro_name}**, and click **Run**.
    6. Happy with the result? Save with **File > Save As > Excel Macro-Enabled Workbook (.xlsm)** to keep the macro.

    The macro expects your data on a sheet called **{expected_sheet}**.

    **Nothing happened after importing?** That is normal. Importing only adds the macro. Step 5 runs it.

    **Saw an error in Excel?** Copy the message, paste it into the box below, and click *Update the macro*.

    **Using Excel on the web?** Macros do not run there. Open the file in the Excel desktop app.
                    """
                )

            with tab_doc:
                st.caption(
                    "A guide your whole team can follow: the manual steps, what the macro automates, and anything "
                    "to confirm. Only column names and the macro code are sent to the AI service, never your data rows."
                )
                st.text_input("Process name (optional)", key="process_name",
                              placeholder="Example: Format the Employees sheet")
                existing_doc = st.session_state.get("process_doc")

                if st.button("Refresh process document" if existing_doc else "Create process document", key="make_doc"):
                    if not api_key.strip():
                        st.error("No API key found. Add it in the sidebar first.")
                    else:
                        with st.spinner("Writing the process document..."):
                            try:
                                st.session_state["process_doc"] = generate_process_doc(
                                    st.session_state.get("requests", []),
                                    st.session_state.get("process_name", "").strip(),
                                    main_macro_name(code) or "Not available",
                                    st.session_state.get("sheet_name", "Data"),
                                    st.session_state.get("columns", []),
                                    code, language, api_key.strip(), model.strip(), provider,
                                )
                                st.session_state["process_doc_version"] = version
                            except Exception as exc:
                                friendly, detail = friendly_error(exc)
                                st.error(friendly)
                                with st.expander("Technical details"):
                                    st.code(detail, language="text")

                doc_markdown = st.session_state.get("process_doc")
                if doc_markdown:
                    if st.session_state.get("process_doc_version") != version:
                        st.warning(
                            f"This document describes version {st.session_state.get('process_doc_version')} of the "
                            f"macro, but you are on version {version}. Click Refresh process document."
                        )
                    assumed = count_assumptions(doc_markdown)
                    if assumed:
                        st.info(f"{assumed} step(s) are marked \"Assumed - please confirm\". "
                                "Check section 6 with the process owner before you share this document.")

                    st.markdown(doc_markdown)

                    doc_title = st.session_state.get("process_name", "").strip() or (main_macro_name(code) or "")
                    try:
                        doc_bytes = markdown_to_docx(doc_markdown, doc_title)
                    except ImportError:
                        doc_bytes = None
                        st.caption("Install python-docx (pip install python-docx) to enable the Word download.")

                    g1, g2, g3 = st.columns(3)
                    if doc_bytes:
                        g1.download_button(
                            "Word (.docx)", data=doc_bytes, file_name="ProcessDocument.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True, key="dl_docx",
                        )
                    g2.download_button(
                        "Text (.md)", data=doc_markdown.encode("utf-8"), file_name="ProcessDocument.md",
                        mime="text/markdown", use_container_width=True, key="dl_md",
                    )
                    g3.download_button(
                        "Everything (.zip)",
                        data=build_package(code, doc_bytes, doc_markdown, st.session_state.get("workbook")),
                        file_name="automation_package.zip", mime="application/zip",
                        use_container_width=True, key="dl_zip",
                    )

            with tab_code:
                show_code_box(code)

            st.divider()
            st.markdown("#### Not quite right?")
            change_request = st.text_area(
                "Tell me what to change, or paste the error message you saw in Excel",
                key=f"change_{version}",
                height=100,
                placeholder="Example: Use green instead of red, or paste: Run-time error '9': Subscript out of range",
            )
            if st.button("Update the macro"):
                if uploaded is None:
                    st.warning("Upload your file again so I can update the macro for it.")
                elif not change_request.strip():
                    st.warning("Describe the change, or paste the error message, first.")
                elif not api_key.strip():
                    st.error("No API key found. Add it in the sidebar first.")
                else:
                    if produce("update", change_request.strip()):
                        st.session_state["flash"] = f"Updated to version {st.session_state['version']}."
                        st.rerun()
        else:
            st.info("Your macro, and a simple explanation of what it does, will appear here after you click Create my macro.")


# --------------------------------------------------------------------------- #
# Option 2: Process document writer (stand-alone, no macro needed)
# --------------------------------------------------------------------------- #
def render_document_writer():
    left, right = st.columns(2, gap="large")

    def send_to_macro_assistant(text):
        st.session_state["mode"] = MODE_MACRO
        st.session_state["instructions"] = text
        st.session_state["macro_hint"] = ("Your process description was copied into the box below. "
                                          "Upload your Excel file, then click Create my macro.")

    def make_doc(inputs, previous_doc=None, change_request=None):
        """Create or update the document. Returns True on success."""
        with st.spinner("Writing your process document..."):
            try:
                markdown_text = generate_standalone_doc(
                    **inputs, language=language, api_key=api_key.strip(), model=model.strip(),
                    provider=provider, previous_doc=previous_doc, change_request=change_request,
                )
                if not markdown_text.strip():
                    raise ValueError("The AI returned an empty answer. Please try again.")
            except Exception as exc:
                friendly, detail = friendly_error(exc)
                st.error(friendly)
                with st.expander("Technical details"):
                    st.code(detail, language="text")
                return False
        st.session_state["pd_markdown"] = markdown_text
        st.session_state["pd_inputs"] = inputs
        st.session_state["pd_version"] = st.session_state.get("pd_version", 0) + 1 if previous_doc else 1
        return True

    with left:
        st.subheader("1. Describe your process")
        notes = st.text_area(
            "Process description",
            key="pd_description",
            height=200,
            placeholder="Paste your notes or a rough draft, or describe the steps in your own words.\n\n"
            "Example: Every month I export the headcount report, copy it into the template, "
            "check the totals against last month, and send it to my manager.",
            label_visibility="collapsed",
        )
        notes_file = st.file_uploader("Or upload your notes (Word, text or Markdown)",
                                      type=["docx", "txt", "md"], key="pd_file")
        file_text = ""
        if notes_file is not None:
            try:
                file_text = read_notes_file(notes_file)
                with st.expander(f"Text read from {notes_file.name}"):
                    st.text(file_text[:3000] + ("\n..." if len(file_text) > 3000 else ""))
            except ImportError:
                st.error("Reading Word files needs python-docx. Install it with: pip install python-docx")
            except Exception as exc:
                st.error(f"Could not read {notes_file.name}: {exc}")

        st.text_input("Process name (optional)", key="pd_name", placeholder="Example: Monthly headcount report")

        with_macro = st.radio("Does a macro already automate this process?", ["No", "Yes"],
                              horizontal=True, key="pd_has_macro") == "Yes"
        macro_name, macro_code = "", ""
        if with_macro:
            macro_name = st.text_input("Macro name (optional)", key="pd_macro_name",
                                       placeholder="Example: FormatDataSheet").strip()
            macro_file = st.file_uploader("Macro file (.bas), optional", type=["bas", "txt", "vba"], key="pd_macro_file")
            pasted = st.text_area("Or paste the macro code (optional)", key="pd_macro_code", height=110)
            macro_code = pasted.strip()
            if not macro_code and macro_file is not None:
                macro_code = macro_file.getvalue().decode("cp1252", errors="replace").strip()
            if macro_code and not macro_name:
                macro_name = main_macro_name(macro_code) or ""
            st.caption("With the macro code, the document can explain exactly what the macro does.")

        st.caption("Everything you type or upload here is sent to the AI service. "
                   "Leave out names, salaries and other confidential details.")
        create = st.button("Create process document", type="primary", key="pd_create")

    with right:
        st.subheader("2. Your process document")

        if create:
            description = "\n\n".join(t.strip() for t in (notes, file_text) if t.strip())
            if not description:
                st.warning("Describe your process or upload your notes first.")
            elif not api_key.strip():
                st.error("No API key found, so nothing was created. "
                         "Paste it in Connection settings, or use Streamlit secrets or an environment variable.")
            else:
                make_doc(dict(description=description, process_name=st.session_state.get("pd_name", "").strip(),
                              with_macro=with_macro, macro_name=macro_name, macro_code=macro_code))

        document_text = st.session_state.get("pd_markdown")
        if not document_text:
            st.info("Your process document will appear here after you click Create process document.")
            return

        inputs = st.session_state["pd_inputs"]
        version = st.session_state.get("pd_version", 1)

        flash = st.session_state.pop("pd_flash", None)
        if flash:
            st.success(flash)
        st.caption(f"Version {version}")

        assumed = count_assumptions(document_text)
        if assumed:
            st.info(f"{assumed} item(s) are marked \"Assumed - please confirm\". "
                    "Check the notes section with the process owner before you share this document.")

        st.markdown(document_text)

        try:
            word_bytes = markdown_to_docx(document_text, inputs.get("process_name") or "Process document")
        except ImportError:
            word_bytes = None
            st.caption("Install python-docx (pip install python-docx) to enable the Word download.")

        d1, d2 = st.columns(2)
        if word_bytes:
            d1.download_button(
                "Word (.docx)", data=word_bytes, file_name="ProcessDocument.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True, type="primary", key="pd_dl_docx",
            )
        d2.download_button("Text (.md)", data=document_text.encode("utf-8"), file_name="ProcessDocument.md",
                           mime="text/markdown", use_container_width=True, key="pd_dl_md")

        st.divider()
        st.markdown("#### Not quite right?")
        change_request = st.text_area(
            "Tell me what to change",
            key=f"pd_change_{version}",
            height=90,
            placeholder="Example: Add a step for checking the totals, or make it shorter.",
        )
        if st.button("Update the document", key="pd_update"):
            if not api_key.strip():
                st.error("No API key found. Add it in the sidebar first.")
            elif not change_request.strip():
                st.warning("Describe the change first.")
            elif make_doc(inputs, previous_doc=document_text, change_request=change_request.strip()):
                st.session_state["pd_flash"] = f"Updated to version {st.session_state['pd_version']}."
                st.rerun()

        if not inputs.get("with_macro"):
            st.divider()
            st.markdown("#### Want to automate this process?")
            st.caption("Turn your description into an Excel macro request. You will need your Excel file.")
            st.button("Turn this into an Excel macro", key="pd_to_macro", on_click=send_to_macro_assistant,
                      args=(f"Automate this process in Excel:\n{inputs['description'][:1500]}",))


# --------------------------------------------------------------------------- #
# Run the option the user chose
# --------------------------------------------------------------------------- #
if mode == MODE_MACRO:
    render_macro_assistant()
else:
    render_document_writer()
